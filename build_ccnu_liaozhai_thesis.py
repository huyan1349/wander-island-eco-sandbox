from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path("/Users/huyan/.codex/attachments/c52f7d71-fa08-423e-84b2-3045290e7ec0/pasted-text.txt")
OUT_DIR = Path("/Users/huyan/Desktop/wander-island_-eco-sandbox")


EN_TITLE = "Dream, Correction and Compromise: The Narrative Interaction between “Obsession” and “Strangeness” in Liaozhai Zhiyi"
EN_SUBTITLE = "A Study of Shuchi and Huang Ying"
EN_ABSTRACT = (
    "This thesis discusses Shuchi and Huang Ying in Liaozhai Zhiyi, with special attention to the relation between "
    "“obsession” and “strangeness”. In many studies, the obsessed figures in Liaozhai Zhiyi are often understood as "
    "people with true feelings and pure temperament, while supernatural women, flower spirits and other strange figures "
    "are usually seen as compensation for what cannot be gained in real life. However, the two stories show a more "
    "complicated situation. In Shuchi, Lang Yuzhu’s love for books is not simply pure. His belief in “Yan Ruyu” and "
    "“golden millet” also shows his distance from real life and his dependence on the imagination produced by the "
    "imperial examination culture. In Huang Ying, Ma Zicai’s love for chrysanthemums contains a similar problem. "
    "He values elegance, but he is unwilling to face the fact that elegant life also needs material support.\n\n"
    "The strange figures in the two stories do not only fulfill the dreams of the protagonists. Yan Ruyu comes out of "
    "a book, but she asks Lang Yuzhu to stop reading all the time and to enter the world of communication, examination "
    "and human affairs. Tao Sheng and Huang Ying are chrysanthemum spirits, but they do not agree with Ma Zicai’s idea "
    "that selling flowers must be vulgar. Instead, they prove that making a living by one’s own labor is not shameful. "
    "In this sense, “strangeness” becomes a force of correction. It brings the obsessed characters back from their "
    "private dreams to the real world.\n\n"
    "The endings of the two stories are not exactly the same. Shuchi ends with the burning of books, success in the "
    "examination, revenge and withdrawal, so its return to reality is painful and uneasy. Huang Ying ends in a milder "
    "way: the business of chrysanthemums is finally stopped, and wealth is brought back into a scholar’s family order. "
    "Through these two stories, this thesis argues that the fantasy in Liaozhai Zhiyi is not only a way to escape from "
    "reality. It also shows Pu Songling’s clear awareness of the gap between literati ideals and the social order of his time."
)
EN_KEYWORDS = "Liaozhai Zhiyi; Shuchi; Huang Ying; obsession; strangeness"


REFERENCE_LINES = [
    "[1] 鲁迅. 中国小说史略[M]. 北京: 人民文学出版社, 1973.",
    "[2] 刘云汉. 浊世红尘中的纯真世界: 透视《聊斋志异》的痴人形象[J]. 名作欣赏, 2007(5).",
    "[3] 陈葆文. 聊斋志异痴狂士人类型析论[M]. 台北: 里仁书局, 2005.",
    "[4] 余习惠. “痴”:《聊斋志异》形象塑造的美学追求[J]. 株洲工学院学报, 2006, 20(1):78-80.",
    "[5] Zeitlin, Judith T. Historian of the Strange: Pu Songling and the Chinese Classical Tale[M]. Stanford: Stanford University Press, 1993.",
    "[6] 夏中权, 李秋花. 离经叛道·惊世骇俗: 谈《聊斋志异·黄英》的反传统观念[J]. 名作欣赏.",
    "[7] 赵海霞. 论蒲松龄的儒者情怀: 以小说《聊斋志异·黄英》为中心的考察[J]. 西北农林科技大学学报(社会科学版), 2009, 9(6):138-141.",
    "[8] 蒲松龄. 聊斋志异[M]. 于天池, 注; 孙通海, 于天池, 等译. 北京: 中华书局, 2015.",
    "[9] Clunas, Craig. Superfluous Things: Material Culture and Social Status in Early Modern China[M]. Cambridge: Polity Press, 1991.",
]


TOC_ENTRIES = [
    ("一、绪论", 0),
    ("（一）研究缘起", 1),
    ("（二）文献综述", 1),
    ("（三）研究对象与问题", 1),
    ("（四）研究方法", 1),
    ("二、痴：对现实秩序的偏离", 0),
    ("（一）郎玉柱：把书本作为现实", 1),
    ("（二）马子才：清雅姿态与生计的疏离", 1),
    ("（三）两种痴的对照", 1),
    ("三、异：作为纠偏者的神女与花妖", 0),
    ("（一）颜如玉", 1),
    ("（二）陶生与黄英", 1),
    ("（三）“异”的双重身份", 1),
    ("四、结局：科举收编与雅俗合流", 0),
    ("（一）《书痴》的结局：焚书、复仇与刺贪刺虐", 1),
    ("（二）《黄英》的结局：由商返儒与雅俗合流", 1),
    ("（三）对照：蒲松龄的怜惜与分寸", 1),
    ("五、结语", 0),
    ("参考文献", 0),
    ("致 谢", 0),
]


def set_run_font(run, size=None, bold=None, name="宋体", ascii_name="Times New Roman"):
    run.font.name = ascii_name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), ascii_name)
    rfonts.set(qn("w:hAnsi"), ascii_name)


def set_paragraph_font(paragraph, size=12, bold=None, name="宋体", ascii_name="Times New Roman"):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, name=name, ascii_name=ascii_name)


def set_doc_defaults(doc: Document):
    sect = doc.sections[0]
    sect.page_width = Cm(21)
    sect.page_height = Cm(29.7)
    # CCNU samples only state A4 + 1.5 line spacing; keep conventional thesis margins.
    sect.top_margin = Cm(2.5)
    sect.bottom_margin = Cm(2.5)
    sect.left_margin = Cm(2.5)
    sect.right_margin = Cm(2.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.first_line_indent = Pt(24)  # 小四约两字符
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.bold = True
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.first_line_indent = None
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(3)
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 2"].font.size = Pt(14)
    styles["Heading 3"].font.size = Pt(12)

    if "Thesis Quote" not in styles:
        quote = styles.add_style("Thesis Quote", WD_STYLE_TYPE.PARAGRAPH)
        quote.base_style = normal
        quote.font.size = Pt(12)
        if quote._element.rPr is None:
            quote._element.append(OxmlElement("w:rPr"))
        if quote._element.rPr.rFonts is None:
            quote._element.rPr.append(OxmlElement("w:rFonts"))
        quote._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        quote._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        quote._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        quote.paragraph_format.left_indent = Cm(0.74)
        quote.paragraph_format.right_indent = Cm(0)
        quote.paragraph_format.first_line_indent = None
        quote.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_sep)
    run._r.append(fld_char_end)
    return run


def set_page_number_start(section, start=1):
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def add_footer_page_number(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE")
    set_paragraph_font(p, size=10.5)


def add_footer_roman_page_number(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE  \\* ROMAN")
    set_paragraph_font(p, size=10.5)


def add_page_break(doc):
    p = doc.paragraphs[-1] if doc.paragraphs else doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def clean_text(text: str) -> str:
    replacements = {
        '"': "”",
        "“。”": "“",
        "”。蒲翁": "蒲翁",
        "尝尝": "常常",
        "本即以": "本文即以",
        "并为因此": "并未因此",
        "应证": "印证",
        "曾今": "曾经",
        "权利": "权力",
        "淘生": "陶生",
        "冯子才": "马子才",
        "冯才子": "马子才",
        "马才子": "马子才",
        "冯玉柱": "郎玉柱",
        "不请自来": "凭空而来",
        "现实越理想": "现实与理想",
        "看起离奇": "看似离奇",
        "中得以": "得以",
        "带考据": "考据",
        "一但": "一旦",
        "出处的窘迫": "处境的窘迫",
        "责任为": "则认为",
        "一位成全": "一味成全",
        "确实士人": "却是士人",
        "不能推成": "不能推演成",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = text.replace("”", "\"")
    text = text.replace("“", "\"")
    text = text.replace("’", "'")
    text = text.replace("‘", "'")
    text = text.replace("」。", "”。")
    text = text.replace("。”[", "”[")
    text = text.replace("。“[", "”[")
    text = text.replace("“", "\"")
    text = re.sub(r"\s+", " ", text).strip()
    text = text.replace(" 。", "。").replace(" ，", "，")
    return text


def normalize_quotes(text: str) -> str:
    # Keep Chinese thesis punctuation sober: turn straight quotes into Chinese-style paired quotes.
    chars = []
    open_quote = True
    for ch in text:
        if ch == '"':
            chars.append("“" if open_quote else "”")
            open_quote = not open_quote
        else:
            chars.append(ch)
    return "".join(chars)


def source_parts():
    raw = SOURCE.read_text(encoding="utf-8")
    raw = raw.replace("## ——以《书痴》《黄英》为中心", "")
    raw = re.sub(r"^> 学术语体全文稿.*\n", "", raw, flags=re.M)
    abstract = re.search(r"## 摘要\s+(.+?)\s+\*\*关键词\*\*：(.+?)\s+## 绪论", raw, flags=re.S)
    if not abstract:
        raise RuntimeError("Cannot parse Chinese abstract")
    cn_abs = normalize_quotes(clean_text(abstract.group(1)))
    cn_keys = clean_text(abstract.group(2))
    body = raw.split("## 绪论", 1)[1].split("## 参考文献", 1)[0]
    body = "## 绪论" + body
    return cn_abs, cn_keys, body


def map_heading(line: str):
    stripped = line.strip()
    mapping = {
        "## 绪论": ("一、绪论", 1),
        "### 一、研究缘起": ("（一）研究缘起", 2),
        "### 二、文献综述": ("（二）文献综述", 2),
        "### 三、研究对象与问题": ("（三）研究对象与问题", 2),
        "### 四、研究方法": ("（四）研究方法", 2),
        "## 第一章 痴：对现实秩序的偏离": ("二、痴：对现实秩序的偏离", 1),
        "第一节 郎玉柱：把书本作为现实": ("（一）郎玉柱：把书本作为现实", 2),
        "### 第二节 马子才：清雅姿态与生计的疏离": ("（二）马子才：清雅姿态与生计的疏离", 2),
        "### 第三节 两种痴的对照": ("（三）两种痴的对照", 2),
        "## 第二章 异：作为纠偏者的神女与花妖": ("三、异：作为纠偏者的神女与花妖", 1),
        "### 第一节 颜如玉": ("（一）颜如玉", 2),
        "### 第二节 陶生与黄英": ("（二）陶生与黄英", 2),
        "### 第三节 \"异\"的双重身份": ("（三）“异”的双重身份", 2),
        "## 第三章 结局：科举收编与雅俗合流": ("四、结局：科举收编与雅俗合流", 1),
        "### 第一节 《书痴》的结局：焚书、复仇与刺贪刺虐": ("（一）《书痴》的结局：焚书、复仇与刺贪刺虐", 2),
        "### 第二节 《黄英》的结局：由商返儒与雅俗合流": ("（二）《黄英》的结局：由商返儒与雅俗合流", 2),
        "### 第三节 对照：蒲松龄的怜惜与分寸": ("（三）对照：蒲松龄的怜惜与分寸", 2),
        "## 结语": ("五、结语", 1),
    }
    return mapping.get(stripped)


def add_body_paragraph(doc, text, style=None):
    text = normalize_quotes(clean_text(text))
    if not text:
        return None
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if style == "Thesis Quote":
        p.paragraph_format.first_line_indent = None
    else:
        p.paragraph_format.first_line_indent = Pt(24)
    run = p.add_run(text)
    set_run_font(run, size=12)
    return p


def add_centered(doc, text, size=12, bold=False, space_after=0, name="宋体", line_spacing=1.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, name=name)
    return p


def add_labeled_para(doc, label, content):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    r = p.add_run(label)
    set_run_font(r, size=12, bold=True)
    r2 = p.add_run(content)
    set_run_font(r2, size=12)
    return p


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    r = p.add_run("分类号（宋体小三加黑）          论文选题类型              \nU D C                       编号                        ")
    set_run_font(r, size=15, bold=True)
    for _ in range(3):
        doc.add_paragraph()
    add_centered(doc, "本科毕业论文（设计）", size=26, bold=True, name="黑体", space_after=8, line_spacing=1.1)
    add_centered(doc, "题 目", size=18, bold=True, line_spacing=1.0)
    add_centered(doc, "幻梦、纠偏与妥协：《聊斋志异》中“痴”与“异”的叙事互动", size=18, bold=True, line_spacing=1.1)
    add_centered(doc, "——以《书痴》《黄英》为中心", size=18, bold=True, line_spacing=1.1)
    for _ in range(1):
        doc.add_paragraph()
    fields = [
        "学    院                    ",
        "专    业                    ",
        "年    级                    ",
        "学生姓名                    ",
        "学    号                    ",
        "指导教师                    ",
    ]
    for field in fields:
        add_centered(doc, field, size=15, bold=True, line_spacing=1.1)
    add_centered(doc, "二〇    年    月", size=16, bold=True, line_spacing=1.1)
    add_page_break(doc)


def add_declaration(doc):
    add_centered(doc, "华中师范大学", size=16, bold=True)
    add_centered(doc, "学位论文原创性声明", size=16, bold=True, space_after=12)
    add_body_paragraph(
        doc,
        "本人郑重声明：所呈交的学位论文是本人在导师指导下独立进行研究工作所取得的研究成果。除了文中特别加以标注引用的内容外，本论文不包含任何其他个人或集体已经发表或撰写的成果作品。本人完全意识到本声明的法律后果由本人承担。"
    )
    add_body_paragraph(doc, "学位论文作者签名：              日期：      年   月   日")
    doc.add_paragraph()
    add_centered(doc, "学位论文版权使用授权书", size=16, bold=True, space_after=12)
    add_body_paragraph(
        doc,
        "本学位论文作者完全了解学校有关保障、使用学位论文的规定，同意学校保留并向有关学位论文管理部门或机构送交论文的复印件和电子版，允许论文被查阅和借阅。本人授权省级优秀学士学位论文评选机构将本学位论文的全部或部分内容编入有关数据库进行检索，可以采用影印、缩印或扫描等复制手段保存和汇编本学位论文。"
    )
    add_body_paragraph(doc, "本学位论文属于")
    add_body_paragraph(doc, "1、保密  □ ，在_____年解密后适用本授权书。")
    add_body_paragraph(doc, "2、不保密  □。")
    add_body_paragraph(doc, "（请在以上相应方框内打“√”）")
    add_body_paragraph(doc, "学位论文作者签名：              日期：       年   月   日")
    add_body_paragraph(doc, "导师签名：                      日期：       年   月   日")
    add_page_break(doc)


def add_toc(doc, page_map):
    add_centered(doc, "目　录", size=18, bold=True, space_after=8)
    for title, level in TOC_ENTRIES:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.left_indent = Cm(0.74 if level else 0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = p.add_run(title)
        set_run_font(run, size=12)
        p.add_run("\t")
        page = page_map.get(title, "")
        r2 = p.add_run(str(page))
        set_run_font(r2, size=12)


def add_abstracts(doc, cn_abs, cn_keys):
    add_labeled_para(doc, "内容摘要：", cn_abs)
    add_labeled_para(doc, "关 键 词：", cn_keys)
    doc.add_paragraph()
    add_labeled_para(doc, "Title：", EN_TITLE)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    r = p.add_run("——" + EN_SUBTITLE)
    set_run_font(r, size=12)
    for idx, para in enumerate(EN_ABSTRACT.split("\n\n")):
        if idx == 0:
            add_labeled_para(doc, "Abstract：", para)
        else:
            add_body_paragraph(doc, para)
    add_labeled_para(doc, "Key words：", EN_KEYWORDS)
    add_page_break(doc)


def add_title_and_cn_abstract(doc, cn_abs, cn_keys):
    add_centered(doc, "幻梦、纠偏与妥协：《聊斋志异》中“痴”与“异”的叙事互动", size=18, bold=True, line_spacing=1.5)
    add_centered(doc, "——以《书痴》《黄英》为中心", size=15, bold=True, line_spacing=1.5)
    doc.add_paragraph()
    add_centered(doc, "中国语言文学类专业", size=12, bold=True, line_spacing=1.5)
    doc.add_paragraph()
    add_centered(doc, "本科生：                  指导教师：", size=12, bold=True, line_spacing=1.5)
    doc.add_paragraph()
    add_labeled_para(doc, "摘要：", cn_abs)
    add_labeled_para(doc, "关键词：", cn_keys)
    add_page_break(doc)


def add_en_abstract_only(doc):
    add_centered(doc, EN_TITLE, size=15, bold=True, line_spacing=1.5)
    add_centered(doc, "——" + EN_SUBTITLE, size=12, bold=True, line_spacing=1.5)
    doc.add_paragraph()
    for idx, para in enumerate(EN_ABSTRACT.split("\n\n")):
        if idx == 0:
            add_labeled_para(doc, "Abstract：", para)
        else:
            add_body_paragraph(doc, para)
    add_labeled_para(doc, "Key words：", EN_KEYWORDS)
    add_page_break(doc)


def add_body(doc, body_text):
    for raw_line in body_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line in {'"', "“", "”", "'"}:
            continue
        heading = map_heading(line)
        if heading:
            text, level = heading
            p = doc.add_paragraph(text, style=f"Heading {level}")
            p.paragraph_format.keep_with_next = True
            set_paragraph_font(p, size=15 if level == 1 else 14, bold=True)
            continue
        if line.startswith(">"):
            add_body_paragraph(doc, line.lstrip("> ").strip(), style="Thesis Quote")
        elif line.startswith("#"):
            continue
        else:
            add_body_paragraph(doc, line)


def add_references(doc):
    p = doc.add_paragraph("参考文献", style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_font(p, size=14, bold=True)
    for ref in REFERENCE_LINES:
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = None
        para.paragraph_format.left_indent = Cm(0.74)
        para.paragraph_format.first_line_indent = Cm(-0.74)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = para.add_run(ref)
        set_run_font(run, size=12)


def add_ack(doc):
    p = doc.add_paragraph("致    谢", style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_font(p, size=14, bold=True)
    add_body_paragraph(
        doc,
        "本文在选题、资料搜集与写作修改过程中得到了老师和同学的帮助，在此谨致谢忱。由于本人学识有限，文中疏漏之处，恳请各位老师批评指正。"
    )


def build_docx(out_path: Path, toc_json: Path | None = None):
    page_map = {}
    if toc_json and toc_json.exists():
        page_map = json.loads(toc_json.read_text(encoding="utf-8"))
    cn_abs, cn_keys, body = source_parts()
    doc = Document()
    set_doc_defaults(doc)
    first = doc.sections[0]
    set_page_number_start(first, 1)
    add_footer_roman_page_number(first)
    add_title_and_cn_abstract(doc, cn_abs, cn_keys)
    add_en_abstract_only(doc)
    add_toc(doc, page_map)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    set_page_number_start(section, 1)
    add_footer_page_number(section)
    add_body(doc, body)
    add_references(doc)
    add_ack(doc)
    doc.save(out_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True)
    parser.add_argument("--toc-json")
    args = parser.parse_args()
    build_docx(Path(args.out), Path(args.toc_json) if args.toc_json else None)


if __name__ == "__main__":
    main()
